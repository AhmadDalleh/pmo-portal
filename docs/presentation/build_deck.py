"""Builds docs/presentation/PMO_Portal_Management_Presentation.pptx (8 slides with speaker notes).

Run from the repository root:  python docs/presentation/build_deck.py
Needs: pip install python-pptx. All figures are the seeded baseline (reporting date 2026-09-30, seed 20260930).
"""
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

INK = RGBColor(0x1F, 0x24, 0x2B)
MUTED = RGBColor(0x5B, 0x64, 0x70)
RED = RGBColor(0xC8, 0x10, 0x2E)
LIGHT = RGBColor(0xF3, 0xF4, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x1E, 0x7B, 0x4A)
GREY = RGBColor(0x9A, 0xA3, 0xAE)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    """paras: list of (text, size, bold, color) tuples; one paragraph each."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, (t, size, bold, color) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        run = p.add_run()
        run.text = t
        run.font.size, run.font.bold, run.font.color.rgb, run.font.name = Pt(size), bold, color, "Calibri"
    return box


def bullets(slide, x, y, w, h, items, size=17):
    return text(slide, x, y, w, h, [("•  " + t, size, False, INK) for t in items])


def card(slide, x, y, w, h, title, body, fill=LIGHT, title_color=RED):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = 0.06
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    text(slide, x + 0.12, y + 0.08, w - 0.24, 0.4, [(title, 15, True, title_color)])
    text(slide, x + 0.12, y + 0.5, w - 0.24, h - 0.55, [(b, 15, False, INK) for b in body])


def box(slide, x, y, w, h, label, fill=INK, color=WHITE, size=14, label_size=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = 0.12
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.08)
    for i, line in enumerate(label.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size if i == 0 else (label_size or size - 2)), i == 0, color, "Calibri"
    return shape


def arrow(slide, x, y, w, h=0.35):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = GREY
    a.line.fill.background()


def new_slide(title, number, notes):
    s = prs.slides.add_slide(BLANK)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.25), prs.slide_height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = RED
    bar.line.fill.background()
    text(s, 0.6, 0.35, 12.2, 0.9, [(title, 30, True, INK)])
    text(s, 0.6, 7.02, 9, 0.3, [("PMO Portfolio & Financial Control Portal  |  baseline 30 September 2026, all data fictional", 10, False, MUTED)])
    text(s, 12.0, 7.02, 0.8, 0.3, [(str(number), 10, False, MUTED)])
    s.notes_slide.notes_text_frame.text = notes
    return s


# 1 ---------------------------------------------------------------------------------------------------------------
s = new_slide("One trusted view of delivery, cost and risk", 1,
 "Business value. Today the PMO runs on spreadsheets, emails and slide decks, so leadership cannot trust any single number for delivery health, cost exposure or overdue work. "
 "I built one connected portal: operational teams maintain the records, and management gets a separate live dashboard computed from the same SQL database. "
 "The six journeys on the right are exactly what I demonstrate: create a project, update a task and watch the dashboard change, manage milestones and risks, post a cost and see spend recalculate, run a change request through approval with history, and drill from a KPI to the underlying records. "
 "Key message: one database, one set of definitions, enforced on the server.")
text(s, 0.6, 1.5, 5.9, 0.5, [("The problem", 20, True, RED)])
bullets(s, 0.6, 2.0, 5.9, 2.2, ["Projects, expenses, risks and approvals live in spreadsheets, emails and decks", "No single, trusted view of delivery health or cost exposure", "Overdue work and budget overruns are found late"])
text(s, 0.6, 4.3, 5.9, 0.5, [("The outcome", 20, True, RED)])
bullets(s, 0.6, 4.8, 5.9, 2.0, ["One portal for teams to maintain records", "One executive dashboard with live, scope-aware KPIs", "Every number traceable to database rows"])
text(s, 7.0, 1.5, 5.8, 0.5, [("Six journeys, all demonstrable", 20, True, RED)])
for i, j in enumerate(["Project manager creates a project, owner, dates, approved budget",
                       "Team member updates a task; progress and dashboard change",
                       "Manager updates a milestone or risk with validation",
                       "Authorised user posts a cost; spend and utilisation update from SQL",
                       "Change request reviewed in a controlled workflow with history",
                       "Manager filters the dashboard and drills down to the records"]):
    box(s, 7.0, 2.05 + i * 0.78, 0.5, 0.6, str(i + 1), fill=RED, size=16)
    text(s, 7.65, 2.05 + i * 0.78, 5.1, 0.7, [(j, 15, False, INK)], anchor=MSO_ANCHOR.MIDDLE)

# 2 ---------------------------------------------------------------------------------------------------------------
s = new_slide("Solution overview: seven connected modules on one database", 2,
 "The seven modules share users, project ids and business rules: portfolio, tasks and milestones, risks, change requests with approvals, finance, the executive dashboard, and access with audit. "
 "Three roles are enforced on the server: Contributor, Project Manager, Administrator. "
 "The system is not a demo shell: it runs on a generated dataset of 10 projects, 118 tasks, 123 cost transactions and 18 change requests, with the required edge cases present on purpose. "
 "Stack: ASP.NET Core on .NET 10, PostgreSQL 16, Angular with Material and ECharts. Admin screens cover users, audit log, automation monitor and cost categories.")
card(s, 0.6, 1.4, 4.0, 3.2, "Modules", ["Project portfolio", "Tasks and milestones", "Risks and issues", "Change requests and approvals", "Financial management", "Executive dashboard", "Access, activity and admin"])
card(s, 4.8, 1.4, 3.9, 3.2, "Roles (enforced by the API)", ["Contributor: own task progress, submit change requests", "Project Manager: manage assigned projects, post costs, decide requests (never own)", "Administrator: all projects, roles, audit, automation"])
card(s, 8.9, 1.4, 3.85, 3.2, "Stack", ["ASP.NET Core, .NET 10", "PostgreSQL 16, EF Core + Dapper", "Angular 22, Material, ECharts", "JWT + refresh cookie, Serilog"])
text(s, 0.6, 4.7, 12, 0.5, [("Running on a reproducible synthetic dataset (seed 20260930)", 18, True, RED)])
for i, (n, l) in enumerate([("10", "projects"), ("118", "tasks"), ("33", "milestones"), ("27", "risks and issues"), ("18", "change requests"), ("123", "cost transactions")]):
    x = 0.6 + i * 2.05
    box(s, x, 5.25, 1.85, 1.15, n + "\n" + l, fill=LIGHT, color=INK, size=26, label_size=14)
text(s, 0.6, 6.5, 12.2, 0.5, [("Edge cases built in: no-spend and no-forecast projects, overspend, overdue and unassigned tasks, draft and post-cutoff costs, pending and rejected revisions", 12, False, MUTED)])

# 3 ---------------------------------------------------------------------------------------------------------------
s = new_slide("Architecture: thin controllers, rules in services, truth in SQL", 3,
 "Clean Architecture with dependencies pointing inwards: Api, Application, Domain; Infrastructure implements the interfaces. "
 "Controllers only bind the request and map the result. Services own the business rules and transactions. Domain entities hold state transitions, for example posting a cost, and throw a rule exception that one handler maps to an HTTP status and a ProblemDetails body with a correlation id. "
 "Writes use EF Core with optimistic concurrency; the analytics use explicit Dapper SQL with one CTE per one-to-many table so spend is never multiplied by joins. "
 "Authorization has two layers: a role policy on the route and a project-scope check in the service, so out-of-scope ids return 404 and forbidden actions 403, even if the frontend is bypassed. "
 "A background worker handles the outbox. Power BI and other tools read the reporting views, not the CSV files.")
box(s, 0.6, 2.2, 2.4, 1.5, "Angular 22\nsignals, interceptors, no KPI maths", fill=RED)
arrow(s, 3.1, 2.75, 0.55)
for i, (t, d) in enumerate([("Api", "controllers, policies, ProblemDetails"), ("Application", "services, validators, transactions"), ("Domain", "entities and rules")]):
    box(s, 3.8, 1.7 + i * 1.25, 3.6, 1.05, t + "\n" + d, fill=INK, size=15)
arrow(s, 7.5, 2.75, 0.55)
box(s, 8.2, 2.2, 2.4, 1.5, "PostgreSQL 16\nconstraints, triggers, views", fill=INK)
box(s, 8.2, 4.2, 2.4, 1.0, "Outbox worker\nBackgroundService", fill=MUTED, size=14)
box(s, 11.0, 2.2, 1.9, 1.5, "Power BI / psql\nread views", fill=MUTED, size=14)
text(s, 8.2, 3.75, 2.4, 0.4, [("^ same database", 11, False, MUTED)])
bullets(s, 0.6, 5.5, 12.2, 1.4, ["Two-layer authorization: role policy on the route, project scope in the service (404 out of scope, 403 not allowed)",
                                  "Business actions are atomic: entity + history + audit + outbox in one transaction",
                                  "Database is a second line of defence: CHECKs, immutable posted costs, append-only audit"], size=15)

# 4 ---------------------------------------------------------------------------------------------------------------
s = new_slide("Financial management: defensible numbers, controlled changes", 4,
 "Finance is an operational module, not a chart. Budgets: initial approved amount plus approved revisions only; pending and rejected revisions change nothing. "
 "Costs are draft or posted; only posted costs on or before the reporting date count as actual spend, and posted costs are immutable, so a correction is a reversal entry, never a silent edit. "
 "Forecast to complete is the sum of future monthly estimates; EAC is actual plus forecast, and variance is budget minus EAC. A project with no forecast keeps EAC equal to actual and is flagged 'No forecast'; we never invent a forecast. "
 "The chart shows budget, actual and EAC per project from the live database. PRJ-009 is already overspent, PRJ-001 is at 92 percent, PRJ-002 is projected to overrun by 247 thousand. "
 "Controls: only authorised roles post, every mutation records actor and time, and crossing 90 or 100 percent of budget queues an alert through the outbox. Approved change requests never turn into spend by themselves.")
cd = CategoryChartData()
cd.categories = ["001", "002", "003", "004", "005", "006", "007", "008", "009", "010"]
cd.add_series("Approved budget", (1300, 2700, 800, 900, 620, 1000, 1500, 1950, 550, 700))
cd.add_series("Actual spend", (1196, 1647, 0, 738, 248, 380, 675, 1072.5, 599.5, 0))
cd.add_series("EAC", (1286, 2947, 500, 738, 248, 580, 1375, 1772.5, 659.5, 600))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.4), Inches(7.6), Inches(4.6), cd)
ch = gf.chart
ch.has_legend, ch.legend.position, ch.legend.include_in_layout = True, XL_LEGEND_POSITION.BOTTOM, False
ch.legend.font.size = Pt(12)
ch.category_axis.tick_labels.font.size = Pt(11)
ch.value_axis.tick_labels.font.size = Pt(11)
ch.value_axis.has_major_gridlines = False
ch.has_title = True
ch.chart_title.text_frame.text = "Budget vs actual vs EAC by project (AED thousand, PRJ-xxx)"
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
for ser, col in zip(ch.plots[0].series, (GREY, RED, INK)):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = col
card(s, 8.4, 1.4, 4.4, 2.3, "Definitions (applied everywhere)", ["Budget = initial + approved revisions", "Actual = posted costs up to the reporting date", "EAC = actual + forecast to complete", "Variance = budget - EAC"])
card(s, 8.4, 3.85, 4.4, 2.15, "Controls and alerts", ["Draft vs posted; posted is immutable, reversal only", "Actor and time on every mutation", "90% and 100% budget alerts via outbox"])
text(s, 0.6, 6.2, 12.2, 0.7, [("Portfolio at 30 Sep 2026: budget AED 12.02M, spend 6.56M (54.5%), EAC 10.71M. PRJ-009 overspent, PRJ-001 at 92%, PRJ-002 projected to overrun.", 14, True, INK)])

# 5 ---------------------------------------------------------------------------------------------------------------
s = new_slide("Web intelligence: every KPI recomputed from SQL, never patched", 5,
 "The dashboard is a real application component set, not embedded Power BI. KPI cards, charts and drill-downs call backend analytics endpoints that run SQL on each request, scoped to what the signed-in user may see. "
 "Filters: department, project, manager, lifecycle status, timeframe and cost category; they live in the URL, so a view can be shared and survives reload. "
 "Drill-downs open the underlying records: the actual-spend card opens the posted cost transactions whose total equals the card, and overdue or risk visuals open the matching tasks or risks. A test asserts the drill-down total equals the KPI for every filter. "
 "Integrity demonstration: update a task's progress or post a new cost, and the dashboard shows new values on the next load, because nothing is computed in the browser. "
 "Two KPIs to explain: Actual spend is the sum of posted cost transactions up to the reporting date, from cost_transactions, filtered by the selected projects; Portfolio progress is the task-weighted average of progress, from tasks, excluding cancelled tasks, never an average of project averages.")
for i, (v, l) in enumerate([("AED 12.02M", "approved budget"), ("AED 6.56M", "actual spend"), ("54.5%", "utilisation"), ("AED 10.71M", "forecast EAC"), ("8", "overdue tasks"), ("4", "open high-exposure risks")]):
    box(s, 0.6 + (i % 3) * 2.35, 1.5 + (i // 3) * 1.45, 2.2, 1.25, v + "\n" + l, fill=LIGHT, color=INK, size=22, label_size=14)
text(s, 0.6, 4.5, 7.0, 0.4, [("Baseline values, administrator scope, unfiltered", 11, False, MUTED)])
bullets(s, 0.6, 4.95, 7.0, 2.0, ["Six filters, kept in the URL (shareable, survives reload)", "Visuals: project health, budget vs actual vs EAC, monthly trend, spend by category, overdue by owner, risk matrix", "Drill-down: KPI to posted transactions, tasks or risks"], size=15)
card(s, 8.0, 1.5, 4.8, 2.5, "Source of every number", ["Actual spend: posted cost_transactions up to the reporting date", "Progress: task weight x progress, cancelled excluded", "One CTE per one-to-many table, so no join inflates spend"])
card(s, 8.0, 4.2, 4.8, 2.5, "Integrity check", ["Update task progress or post a cost, then reload: KPIs change", "No client-side KPI maths", "Test: drill-down rows sum to the KPI, for every filter"])

# 6 ---------------------------------------------------------------------------------------------------------------
s = new_slide("Automation: reliable notifications through a transactional outbox", 6,
 "The automation turns business events into notifications without ever risking the business transaction. "
 "Trigger: approving a change request, or a posted cost taking a project past 90 or 100 percent of its budget. The same database transaction that makes the change also writes an outbox row. "
 "A background worker polls every ten seconds, claims pending rows with SKIP LOCKED so two workers never take the same row, creates in-app notifications, and records an automation run with counts and a correlation id. "
 "Safety: a unique dedup key stops the same event being queued twice, a unique constraint on message and user stops a second notification, failures retry with exponential backoff and end as 'failed' after five attempts with the error text, and an administrator can retry from the Automation screen. "
 "Evidence: I reset the database, the worker processed the pending message by itself, I retried the seeded failed message from the API and it was processed on the next run. The write-up is in docs/automation-evidence.md.")
for i, (t, d) in enumerate([("Event", "CR approved or\nbudget 90% / 100%"), ("Outbox row", "same transaction\nas the change"), ("Worker", "every 10 s, SKIP LOCKED"), ("Notification", "per user, no duplicates"), ("Run log", "counts, error,\ncorrelation id")]):
    x = 0.6 + i * 2.5
    box(s, x, 1.7, 2.1, 1.3, t + "\n" + d, fill=INK if i != 1 else RED, size=15)
    if i < 4:
        arrow(s, x + 2.12, 2.2, 0.35, 0.3)
card(s, 0.6, 3.4, 6.0, 2.6, "Safe by design", ["Same event cannot be queued twice (unique dedup key)", "Same message cannot notify a user twice", "Failure: backoff, then 'failed' after 5 attempts, with last error", "Admin Retry (with confirm) puts it back in the queue"])
card(s, 6.85, 3.4, 5.9, 2.6, "Evidence (docs/automation-evidence.md)", ["Fresh reset: message 9 processed by the worker (run recorded)", "Failed message 7 retried via API, processed on next run", "3 notifications, no duplicates", "Automated tests cover duplicate, failure and retry paths"])
text(s, 0.6, 6.25, 12.2, 0.6, [("Message 7 is a seeded 'simulated' failure, not a defect: it exists to demonstrate the retry path.", 13, False, MUTED)])

# 7 ---------------------------------------------------------------------------------------------------------------
s = new_slide("AI approach: specs first, tests as the judge, ownership throughout", 7,
 "AI was a development accelerator, but I stayed the owner of every rule. Strategy: I wrote the requirements and financial definitions in documents first; AI worked slice by slice from story ids, and every slice ended with an explanation I had to be able to give back. "
 "Business rules and formulas come from my design document; I verify them with tests against the seeded edge cases rather than trusting the generated code. For debugging I gave the AI the failing test, the log line with the correlation id and the SQL, and asked for a hypothesis before a fix. "
 "The log has eight entries. Three examples where the AI was wrong: integer division made the portfolio progress KPI 63 instead of 63.24, found because a test computes the expected value with its own SQL; the seed data used different audit action names than the API; and a shared table change compiled but crashed at runtime, found in the browser console. "
 "I will open any file, trace a request end to end and make a live change. [Presenter: adjust the last line to what you personally validated.]")
card(s, 0.6, 1.5, 6.0, 2.5, "Prompt strategy", ["Specs first: stories and formulas in docs, prompts cite story ids", "Plan, approve, implement, explain-back per vertical slice", "Rules written by me; AI implements, tests judge", "Debugging: failing test + log + SQL, hypothesis before fix"])
card(s, 6.85, 1.5, 5.9, 2.5, "Where the AI was wrong (and how I caught it)", ["Integer division: progress 63 vs 63.24 (test with independent SQL)", "Seed audit names differed from the API (endpoint test)", "Green build, broken table at runtime (browser console)", "Wrong claim about the importer (read the code, re-ran it)"], title_color=RED)
card(s, 0.6, 4.25, 12.15, 2.4, "Ownership", ["8 log entries with problem, prompt, result, accepted or rejected, test evidence; tutor walkthrough per module and a list of where code differs from the docs", "I can trace any request end to end, explain the SQL and make a live change", "Disclosure of AI-assisted areas and what I validated personally is in docs/ai-development-log.md"], fill=LIGHT)

# 8 ---------------------------------------------------------------------------------------------------------------
s = new_slide("Testing and delivery: proven rules, reproducible setup", 8,
 "Testing: 246 automated checks. 51 unit tests for domain rules and financial calculations, and 195 integration tests that run the real API against a real PostgreSQL in a container: workflow transitions, financial figures with exact decimals, concurrency, idempotency and RBAC. "
 "Acceptance test D is automated: a Contributor calling approval or cost posting directly gets 403 and nothing changes. A stale version returns 409, and a repeated Idempotency-Key creates one cost. "
 "Delivery: Git history of small conventional commits, a README with setup, an environment example with no secrets, migrations, a generator with a fixed seed, generated CSVs, an import script and a reset script. "
 "Honest status: the reporting views, report queries, DAX measures and reconciliation template are ready; the Power BI report file, the Access utility and a hosted URL are still to do. [Presenter: update this line before the interview.] "
 "Everything runs locally; the frontend, API and SQL database are verifiable end to end.")
card(s, 0.6, 1.5, 4.0, 3.3, "Automated checks: 246", ["51 unit: rules, EAC, transitions", "195 integration (Testcontainers PostgreSQL)", "RBAC: Contributor denied on approve and post", "409 on stale version, one cost per Idempotency-Key", "Error log with correlation id"])
card(s, 4.8, 1.5, 4.0, 3.3, "Security", ["JWT 15 min, rotating refresh cookie (HttpOnly)", "Role policy + project scope on every route", "Parameterised SQL, allow-listed sorting", "Rate limit and lockout on login", "Secrets in user-secrets / env, never in Git"])
card(s, 9.0, 1.5, 3.75, 3.3, "Delivery", ["Small conventional commits, README setup", ".env.example, EF migrations", "Seeded generator + 19 CSVs", "import-data and reset-db scripts", "Swagger + OpenAPI export"])
card(s, 0.6, 5.0, 12.15, 1.8, "Status and next steps", ["Done: seven modules, dashboard, finance, automation, admin screens, reporting views and SQL reports, Power BI model and DAX, reconciliation template", "To do: Power BI report file, Access utility, hosted URL (optional); browser walkthrough of every flow"], fill=LIGHT)

out = "docs/presentation/PMO_Portal_Management_Presentation.pptx"
prs.save(out)
print("saved", out, len(prs.slides), "slides")
